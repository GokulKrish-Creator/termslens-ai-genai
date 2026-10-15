from io import BytesIO
from urllib.parse import urlparse
import re

import requests
from bs4 import BeautifulSoup


def extract_uploaded_file(filename: str, data: bytes):
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(data))
        pages = []
        for i, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            pages.append(f"[Page {i}]\n{text}")
        return "\n\n".join(pages), filename
    if name.endswith(".docx"):
        from docx import Document
        doc = Document(BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts), filename
    if name.endswith((".html", ".htm")):
        soup = BeautifulSoup(data.decode("utf-8", errors="replace"), "html.parser")
        for tag in soup(["script", "style", "noscript", "nav", "footer"]):
            tag.decompose()
        return clean_text(soup.get_text("\n")), filename
    return data.decode("utf-8", errors="replace"), filename


def clean_text(text: str) -> str:
    text = text.replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n[ \t]*\n[ \t]*\n+", "\n\n", text)
    return text.strip()


def fetch_website(url: str):
    parsed = urlparse((url or "").strip())
    if parsed.scheme not in ("http", "https") or not parsed.netloc:
        raise ValueError("Enter a complete public URL starting with https:// or http://")
    if parsed.username or parsed.password:
        raise ValueError("URLs containing embedded credentials are not supported.")
    response = requests.get(
        url.strip(),
        timeout=20,
        headers={"User-Agent": "TermsLensAI/1.0 (public policy text reader)"},
        allow_redirects=True,
    )
    response.raise_for_status()
    content_type = response.headers.get("content-type", "").lower()
    if "html" not in content_type and not response.text.lstrip().lower().startswith(("<!doctype html", "<html")):
        raise ValueError("The URL did not return an HTML webpage. Download the policy as a file and upload it instead.")
    soup = BeautifulSoup(response.text, "html.parser")
    title = soup.title.get_text(" ", strip=True) if soup.title else parsed.netloc
    for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "svg", "form"]):
        tag.decompose()
    main = soup.find("main") or soup.find("article") or soup.body or soup
    text = clean_text(main.get_text("\n"))
    if len(text) > 180_000:
        text = text[:180_000] + "\n\n[Text truncated by TermsLens AI for processing.]"
    return text, title


def split_into_chunks(text: str, max_chars: int = 12000):
    text = text.strip()
    if not text:
        return []
    chunks = []
    while len(text) > max_chars:
        cut = text.rfind("\n", 0, max_chars)
        if cut < max_chars // 2:
            cut = max_chars
        chunks.append(text[:cut].strip())
        text = text[cut:].strip()
    if text:
        chunks.append(text)
    return chunks
